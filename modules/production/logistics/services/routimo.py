# -*- coding: utf-8 -*-
"""
Eksport do Routimo — wspólny generator Excela (spec 8.4).

Formatowanie przeniesione 1:1 z modules/reports/routers.generate_routimo_excel;
raporty budują swoje wiersze i wołają zbuduj_excel(), trasy logistyki —
wiersze_trasy(). Kolumny: 37, kolejność jak w NAGLOWKI.
"""
import io
import re
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy.orm import selectinload

from modules.production.logistics.adresy import extract_house_and_apartment_number
from modules.production.logistics.services import delivery, geocoding, routes
from modules.production.models import ProductionOrder

KOLUMNA_KOMENTARZA = 33  # AG — wieloliniowa lista produktów

NAGLOWKI = [
    'Nazwa', 'Klient', 'Nazwa przesyłki', 'Numer wew.', 'Koszty kuriera netto', 'Ulica',
    'Numer domu', 'Numer mieszkania', 'Kod pocztowy', 'Miasto', 'Kraj', 'Region',
    'Numer telefonu', 'Email', 'Email klienta', 'Nip klienta', 'Początek okna czasowego',
    'Koniec okna czasowego', 'Okno czasowe', 'Czas na wykonanie zadania',
    'Oczekiwana data realizacji', 'Harmonogram', 'Pojazd', 'Typy pojazdów',
    'Liczba przesyłek', 'Wielkość przesyłki', 'Waga przesyłki', 'Wartość przesyłki',
    'Forma płatności', 'Waluta', 'Szerokość geograficzna', 'Długość geograficzna',
    'Komentarz', 'Komentarz 2', 'Uwagi', 'Dodatkowe 1', 'Dodatkowe 2',
]

SZEROKOSCI = {
    'A': 40.0, 'B': 31.81, 'C': 17.0, 'D': 13.0, 'E': 13.0, 'F': 32.0, 'G': 9.0, 'H': 9.0,
    'I': 14.0, 'J': 25.0, 'K': 12.0, 'L': 20.0, 'M': 15.0, 'N': 25.0, 'O': 38.0, 'P': 15.0,
    'Q': 20.0, 'R': 20.0, 'S': 15.0, 'T': 25.0, 'U': 20.0, 'V': 15.0, 'W': 15.0, 'X': 20.0,
    'Y': 15.0, 'Z': 18.0, 'AA': 15.0, 'AB': 18.0, 'AC': 20.0, 'AD': 10.0, 'AE': 20.0,
    'AF': 20.0, 'AG': 70.0, 'AH': 20.0, 'AI': 25.0, 'AJ': 15.0, 'AK': 15.0,
}

# R10 (kontroler): transliteracja PRZED slugowaniem nazwy pliku — unicodedata NFKD
# NIE rozkłada ł/Ł na osobne znaki (to nie jest znak + akcent, tylko osobna litera),
# więc jawna mapa najpierw, a dopiero potem NFKD + ASCII na resztę ewentualnych
# akcentów (nazwy tras nie muszą być polskie).
_TRANSLITERACJA_PL = {
    'ą': 'a', 'ć': 'c', 'ę': 'e', 'ł': 'l', 'ń': 'n', 'ó': 'o', 'ś': 's', 'ź': 'z', 'ż': 'z',
    'Ą': 'A', 'Ć': 'C', 'Ę': 'E', 'Ł': 'L', 'Ń': 'N', 'Ó': 'O', 'Ś': 'S', 'Ź': 'Z', 'Ż': 'Z',
}


def zbuduj_excel(wiersze):
    """Lista wierszy (każdy = 37 wartości w kolejności NAGLOWKI) → bytes pliku .xlsx."""
    skoroszyt = openpyxl.Workbook()
    arkusz = skoroszyt.active
    arkusz.title = 'Sheet1'
    skoroszyt.create_sheet('Sheet2')  # pusty drugi arkusz, jak we wzorcu Routimo

    wypelnienie = PatternFill(start_color='F3F3F3', end_color='EFEFEF', fill_type='solid')
    czcionka = Font(bold=True, underline='single')
    wyrownanie = Alignment(horizontal='left', vertical='center', wrap_text=True)
    linia = Side(border_style='thin', color='000000')
    obramowanie = Border(left=linia, right=linia, top=linia, bottom=linia)
    for kolumna, naglowek in enumerate(NAGLOWKI, 1):
        komorka = arkusz.cell(row=1, column=kolumna, value=naglowek)
        komorka.fill, komorka.font = wypelnienie, czcionka
        komorka.alignment, komorka.border = wyrownanie, obramowanie
    for litera, szerokosc in SZEROKOSCI.items():
        arkusz.column_dimensions[litera].width = szerokosc
    arkusz.row_dimensions[1].height = 43.0

    for nr_wiersza, wiersz in enumerate(wiersze, 2):
        for kolumna, wartosc in enumerate(wiersz, 1):
            komorka = arkusz.cell(row=nr_wiersza, column=kolumna, value=wartosc)
            if kolumna == KOLUMNA_KOMENTARZA:
                komorka.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        komentarz = wiersz[KOLUMNA_KOMENTARZA - 1] if len(wiersz) >= KOLUMNA_KOMENTARZA else None
        if komentarz and '\n' in str(komentarz):
            linie = str(komentarz).count('\n') + 1
            arkusz.row_dimensions[nr_wiersza].height = max(15 * linie, 15)

    bufor = io.BytesIO()
    skoroszyt.save(bufor)
    return bufor.getvalue()


def _produkty(order):
    return '\n'.join(u'{} x{}'.format(p.original_product_name, p.quantity or 1)
                     for p in delivery.aktywne_produkty(order))


def _zamowienia_trasy_z_produktami(route):
    """
    Zamówienia trasy w kolejności przystanków, z pozycjami (order.products)
    doładowanymi JEDNYM zapytaniem (ruling Task 7: żadnych zapytań per zamówienie).
    `routes.zamowienia_trasy` samo w sobie nie dociąga `products`, a każdy wiersz
    poniżej go czyta (_produkty, m3, wartość) — bez selectinload to N dodatkowych
    zapytań na eksport trasy. Bez konfiguracji produktu (routimo jej nie czyta,
    w przeciwieństwie do trasy_api._zamowienia_z_produktami dla panelu).
    """
    ids = [s.order_id for s in route.stops]
    if not ids:
        return []
    po_id = {o.id: o for o in ProductionOrder.query.options(
        selectinload(ProductionOrder.products)
    ).filter(ProductionOrder.id.in_(ids)).all()}
    return [po_id[i] for i in ids if i in po_id]


def wiersze_trasy(route):
    from modules.reports.utils import PostcodeToStateMapper
    zamowienia = _zamowienia_trasy_z_produktami(route)
    punkty = geocoding.geo_zamowien([o.id for o in zamowienia])
    pojazd = route.vehicle.name if route.vehicle is not None else ''
    wiersze = []
    for order in zamowienia:
        aktywne = delivery.aktywne_produkty(order)
        dom, mieszkanie, ulica = extract_house_and_apartment_number(order.delivery_address or '')
        m3 = sum(float(p.volume_m3 or 0) * (p.quantity or 1) for p in aktywne)
        wartosc = sum(float(p.total_value_net or 0) for p in aktywne)
        punkt = punkty.get(order.id)
        kraj = (order.delivery_country_code or 'PL').upper()
        wiersze.append([
            order.client_name or '', order.client_name or '', order.baselinker_order_id,
            order.internal_order_number or '', '', ulica, dom, mieszkanie,
            order.delivery_postcode or '', order.delivery_city or '',
            'Polska' if kraj == 'PL' else kraj,
            PostcodeToStateMapper.get_state_from_postcode(order.delivery_postcode or '') or '',
            order.client_phone or '', '', order.client_email or '', '', '', '', '', '',
            route.date_from.isoformat(), '', pojazd, '',
            sum(int(p.quantity or 0) for p in aktywne), round(m3, 3),
            round(m3 * routes.WAGA_KG_NA_M3, 2), round(wartosc, 2), '', 'PLN',
            float(punkt.lat) if punkt is not None and punkt.lat is not None else '',
            float(punkt.lng) if punkt is not None and punkt.lng is not None else '',
            _produkty(order),
            u'{}, {}'.format(order.baselinker_order_id or '', order.internal_order_number or '').strip(', '),
            '', '', '',
        ])
    return wiersze


def _bez_polskich_znakow(tekst):
    zamienione = ''.join(_TRANSLITERACJA_PL.get(znak, znak) for znak in tekst)
    return unicodedata.normalize('NFKD', zamienione).encode('ascii', 'ignore').decode('ascii')


def nazwa_pliku(route):
    slug = re.sub(r'[^a-z0-9]+', '-', _bez_polskich_znakow(route.name).lower()).strip('-') or 'trasa'
    return 'routimo_{}_{}.xlsx'.format(slug, route.date_from.isoformat())
